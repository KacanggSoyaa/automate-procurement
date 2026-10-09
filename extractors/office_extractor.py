"""Word (.docx) and Excel (.xlsx/.xlsm) text extraction for ITB packages."""
import fnmatch
import os
from typing import Dict, List

# Office document extensions this extractor can read.
DOCX_EXTENSIONS = ('.docx',)
EXCEL_EXTENSIONS = ('.xlsx', '.xlsm')


class OfficeExtractor:
    """Reads text out of Word and Excel documents found in a folder."""

    def extract_docx_text(self, path: str) -> str:
        """Return paragraphs and tables of a .docx as plain text lines."""
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError(
                "python-docx is required to read .docx files (pip install python-docx)."
            ) from e
        doc = Document(path)
        parts: List[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)

    def extract_xlsx_text(self, path: str) -> str:
        """Return every non-empty spreadsheet row as ``a | b | c`` text lines.

        Each sheet is prefixed with a ``# Sheet: <name>`` marker so the model
        can tell workbook tabs apart.
        """
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError(
                "openpyxl is required to read .xlsx files (pip install openpyxl)."
            ) from e
        wb = load_workbook(path, read_only=True, data_only=True)
        parts: List[str] = []
        try:
            for ws in wb.worksheets:
                parts.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v).strip() for v in row]
                    if any(cells):
                        parts.append(" | ".join(cells))
        finally:
            wb.close()
        return "\n".join(parts)

    def extract_office_folder(self, folder: str, pattern: str = '*') -> List[Dict]:
        """Read every supported Word/Excel file in a folder matching a glob.

        Returns the same shape as ``PDFExtractor.extract_itb_folder``: one dict
        per file with ``filename``, ``path``, ``content`` and ``size``, or an
        ``error`` key with empty content when a file cannot be read.
        """
        results: List[Dict] = []
        for fname in os.listdir(folder):
            low = fname.lower()
            if low.endswith(DOCX_EXTENSIONS):
                reader = self.extract_docx_text
            elif low.endswith(EXCEL_EXTENSIONS):
                reader = self.extract_xlsx_text
            else:
                continue
            if not fnmatch.fnmatch(fname, pattern):
                continue
            path = os.path.join(folder, fname)
            try:
                content = reader(path)
                results.append({
                    'filename': fname,
                    'path': path,
                    'content': content,
                    'size': len(content),
                })
            except Exception as e:
                results.append({
                    'filename': fname,
                    'path': path,
                    'error': str(e),
                    'content': '',
                })
        return results
