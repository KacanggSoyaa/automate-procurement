from docx import Document
from docx.shared import Inches, Pt
from pathlib import Path
import datetime

class SpecSheetGenerator:
    def generate_spec_sheet(self, extracted_data: dict, output_path: str):
        doc = Document()
        # Title
        title = doc.add_heading('Specification Sheet & Reference Document', level=0)
        
        # Metadata
        doc.add_paragraph(f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph("")
        
        # Summary
        doc.add_heading('1. Summary', level=1)
        doc.add_paragraph(f"Source files processed: {extracted_data.get('files_processed', 0)}")
        doc.add_paragraph(f"Line items extracted: {extracted_data.get('line_items_count', 0)}")
        doc.add_paragraph("")
        
        # Technical Requirements
        doc.add_heading('2. Technical Requirements', level=1)
        for req in extracted_data.get('technical_requirements', []):
            doc.add_paragraph(req, style='List Bullet')
        doc.add_paragraph("")
        
        # Line Items
        doc.add_heading('3. Line Items', level=1)
        for item in extracted_data.get('line_items', []):
            p = doc.add_paragraph()
            p.add_run(f"Item {item.get('item_no', '')}: ").bold = True
            p.add_run(item.get('description', ''))
        doc.add_paragraph("")
        
        # Submission Rules
        doc.add_heading('4. Submission Rules', level=1)
        for rule in extracted_data.get('submission_rules', []):
            doc.add_paragraph(rule, style='List Bullet')
        
        doc.save(output_path)
        return output_path
