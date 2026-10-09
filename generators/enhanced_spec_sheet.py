"""Builds the teal RFQ specification sheet DOCX from classified ITB data."""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# Corporate colour palette used throughout the generated document.
PRIMARY_COLOR = RGBColor(0, 128, 128)
SECONDARY_COLOR = RGBColor(51, 51, 51)
MUTED_TEXT = RGBColor(100, 100, 100)
HEADER_BG = "008080"
ALT_ROW_BG = "F4F8F8"


class EnhancedSpecSheetGenerator:
    """Renders classified ITB data into the formatted RFQ specification DOCX."""

    def _set_cell_background(self, cell, hex_color):
        """Fill a table cell with a solid background colour."""
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def _set_cell_margins(self, cell, top=100, bottom=100, left=150, right=150):
        """Apply internal padding (in dxa twips) to a table cell."""
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def _add_heading_styled(self, doc, text, level):
        """Add a level 1/2 heading with the corporate font and colour."""
        h = doc.add_heading(level=level)
        run = h.add_run(text)
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = PRIMARY_COLOR
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
        elif level == 2:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = SECONDARY_COLOR
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
        return h

    def _meta_block(self, doc, meta):
        """Add the project/reference and delivery metadata table."""
        project_ref = meta.get('project_ref') or 'To be specified'
        project_title = meta.get('project_title') or ''
        project_line = f"Project / Ref: {project_ref}"
        if project_title:
            project_line += f" - {project_title}"
        delivery = meta.get('delivery_location') or 'To be specified'

        meta_table = doc.add_table(rows=2, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        meta_data = [
            (project_line, f"Delivery Location: {delivery}"),
            ("Document Type: Technical Specification Sheet", "Target Delivery: To be specified"),
        ]
        widths = [Inches(3.8), Inches(3.0)]

        for row_idx, (col1, col2) in enumerate(meta_data):
            row = meta_table.rows[row_idx]
            for col_idx, text in enumerate((col1, col2)):
                cell = row.cells[col_idx]
                cell.width = widths[col_idx]
                self._set_cell_background(cell, "F0F4F4")
                self._set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                p = cell.paragraphs[0]
                r = p.add_run(text)
                r.font.name = 'Calibri'
                r.font.bold = True
                r.font.size = Pt(9.5)
                p.paragraph_format.space_after = Pt(1)

    def _items_table(self, doc, items):
        """Add the summary table of extracted line items."""
        table = doc.add_table(rows=1 + max(len(items), 1), cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        headers = ["No.", "Item Name", "Item Number", "Description & Key Specs", "Qty", "UOM"]
        col_widths = [Inches(0.5), Inches(1.8), Inches(1.1), Inches(2.5), Inches(0.5), Inches(0.4)]

        hdr_cells = table.rows[0].cells
        for idx, text in enumerate(headers):
            hdr_cells[idx].width = col_widths[idx]
            self._set_cell_background(hdr_cells[idx], HEADER_BG)
            self._set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=80, right=80)
            p = hdr_cells[idx].paragraphs[0]
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

        if not items:
            cells = table.rows[1].cells
            cells[0].merge(cells[5])
            p = cells[0].paragraphs[0]
            r = p.add_run("No line items identified.")
            r.font.name = 'Calibri'
            r.font.size = Pt(9.5)
            return

        for i, item in enumerate(items, 1):
            row_cells = table.rows[i].cells
            desc = item.get('description') or ''
            spec_bits = [f"{s.get('parameter')}: {s.get('value')}" for s in item.get('specs', []) if s.get('parameter')]
            if spec_bits:
                desc = (desc + " | " if desc else "") + "; ".join(spec_bits)
            data = [
                str(item.get('item_no') or i),
                item.get('item_name') or desc,
                item.get('item_number') or '',
                desc,
                str(item.get('qty') or ''),
                item.get('uom') or item.get('unit') or '',
            ]
            for idx, text in enumerate(data):
                row_cells[idx].width = col_widths[idx]
                self._set_cell_background(row_cells[idx], "FFFFFF" if i % 2 else ALT_ROW_BG)
                self._set_cell_margins(row_cells[idx], top=70, bottom=70, left=80, right=80)
                p = row_cells[idx].paragraphs[0]
                r = p.add_run(text)
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                if idx in [0, 2, 4, 5]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _details(self, doc, items):
        """Add a per-item detailed technical specification section."""
        for i, item in enumerate(items, 1):
            title = item.get('item_name') or (item.get('description') or f'Item {i}')
            if len(title) > 80:
                title = title[:77] + '...'
            self._add_heading_styled(doc, f"Item {i}: {title}", level=2)

            rows = []
            if item.get('item_name'):
                rows.append(("Item Name", item['item_name']))
            if item.get('item_number'):
                rows.append(("Item Number", item['item_number']))
            if item.get('description'):
                rows.append(("Description", item['description']))
            for s in item.get('specs', []):
                if s.get('parameter'):
                    rows.append((s['parameter'], s.get('value', '')))
            if item.get('qty'):
                unit = item.get('uom') or item.get('unit') or ''
                rows.append(("Quantity Required", f"{item['qty']} {unit}".strip()))

            if not rows:
                rows.append(("Description", "Details not extracted."))

            for param, val in rows:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Inches(0.2)
                r_p = p.add_run(f"{param}: ")
                r_p.font.name = 'Calibri'
                r_p.font.bold = True
                r_p.font.size = Pt(10)
                r_v = p.add_run(str(val))
                r_v.font.name = 'Calibri'
                r_v.font.size = Pt(10)

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def _instructions(self, doc, submission_rules):
        """Add the quotation submission instructions, falling back to defaults."""
        instructions = submission_rules[:12] if submission_rules else [
            "Please provide unit price and total pricing delivered to the specified delivery location.",
            "Include official manufacturer datasheets or technical compliance verification for the proposed model.",
            "Specify the lead time / estimated delivery date.",
            "State the standard manufacturer warranty period and warranty support terms.",
            "Ensure quotation validity is a minimum of 30 to 60 days from submission date.",
        ]
        for inst in instructions:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(inst)
            r.font.name = 'Calibri'
            r.font.size = Pt(10)

    def _source_note(self, doc, source_files):
        """Add the small evidence footer listing the source documents."""
        if not source_files:
            return
        doc.add_paragraph()
        heading = doc.add_paragraph()
        hr = heading.add_run("Source Documents (evidence):")
        hr.font.name = 'Calibri'
        hr.font.bold = True
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(110, 110, 110)
        heading.paragraph_format.space_after = Pt(2)
        for name in source_files:
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(name)
            r.font.name = 'Calibri'
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(110, 110, 110)
            p.paragraph_format.space_after = Pt(0)

    def generate(self, extracted_data: dict, output_path: str):
        """Generate the full specification sheet and return the output path."""
        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        title_p = doc.add_paragraph()
        title_run = title_p.add_run("REQUEST FOR QUOTATION (RFQ)")
        title_run.font.name = 'Calibri'
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = PRIMARY_COLOR
        title_p.paragraph_format.space_after = Pt(2)

        subtitle_p = doc.add_paragraph()
        sub_run = subtitle_p.add_run("Technical Specifications & Item Requirement Schedule")
        sub_run.font.name = 'Calibri'
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.color.rgb = MUTED_TEXT
        subtitle_p.paragraph_format.space_after = Pt(12)

        items = extracted_data.get('line_items', [])
        meta = extracted_data.get('meta', {}) or {}

        self._meta_block(doc, meta)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        self._add_heading_styled(doc, "1. Summary of Required Items", level=1)
        self._items_table(doc, items)
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        self._add_heading_styled(doc, "2. Detailed Technical Specifications", level=1)
        self._details(doc, items)

        self._add_heading_styled(doc, "3. Quotation Submission Instructions", level=1)
        self._instructions(doc, extracted_data.get('submission_rules', []))

        self._source_note(doc, extracted_data.get("source_files", []))

        doc.save(output_path)
        return output_path
